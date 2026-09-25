"""Build the report .docx from a report model (python-docx).

The model is plain data assembled by the skill (narrative from Opus, numbers from
compute.py, chart PNGs from charts.py). docx_writer only lays it out — it does not
generate content or numbers.

Report model:
  {"title": str, "subtitle": str?, "meta": {date,...}?,
   "sections": [{"heading": str, "level": 1|2, "paragraphs": [str],
                 "bullets": [str]?, "charts": [png_path]?}],
   "methodology": str?,                 # methods/sources box (top-tier provenance)
   "appendix": {"questions": [str], "comments_prompt": str}? }  # drafts 1-2 only; None for final
"""
from __future__ import annotations

from pathlib import Path

NAVY = (0x21, 0x24, 0x30)
PRIMARY = (0x5A, 0x52, 0xFF)


def _rgb(t):
    from docx.shared import RGBColor
    return RGBColor(*t)


def build_report(model, out_path):
    """Render `model` to a .docx at out_path. Returns the Path."""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # --- title block ---
    t = doc.add_paragraph()
    run = t.add_run(model["title"])
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = _rgb(NAVY)
    if model.get("subtitle"):
        sp = doc.add_paragraph()
        sr = sp.add_run(model["subtitle"])
        sr.font.size = Pt(13)
        sr.font.color.rgb = _rgb(PRIMARY)
    if model.get("meta"):
        meta = " · ".join(f"{k}: {v}" for k, v in model["meta"].items())
        mp = doc.add_paragraph()
        mr = mp.add_run(meta)
        mr.italic = True
        mr.font.size = Pt(9)
        mr.font.color.rgb = _rgb(NAVY)
    doc.add_paragraph()

    # --- sections ---
    for sec in model.get("sections", []):
        level = sec.get("level", 1)
        h = doc.add_heading(sec["heading"], level=min(level, 2))
        for r in h.runs:
            r.font.color.rgb = _rgb(NAVY)
        for para in sec.get("paragraphs", []):
            doc.add_paragraph(para)
        for b in sec.get("bullets", []) or []:
            doc.add_paragraph(b, style="List Bullet")
        for png in sec.get("charts", []) or []:
            if Path(png).exists():
                doc.add_picture(str(png), width=Inches(6.0))
                cap = doc.paragraphs[-1]
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- methodology box ---
    if model.get("methodology"):
        doc.add_paragraph()
        mh = doc.add_heading("Metodologia e fonti", level=2)
        for r in mh.runs:
            r.font.color.rgb = _rgb(NAVY)
        mb = doc.add_paragraph()
        mr = mb.add_run(model["methodology"])
        mr.font.size = Pt(9)
        mr.italic = True

    # --- gate appendix (drafts 1-2 only) ---
    app = model.get("appendix")
    if app:
        doc.add_page_break()
        ah = doc.add_heading("Domande per continuare", level=1)
        for r in ah.runs:
            r.font.color.rgb = _rgb(PRIMARY)
        for q in app.get("questions", []):
            doc.add_paragraph(q, style="List Number")
        doc.add_paragraph()
        ch = doc.add_heading("Commenti finali", level=1)
        for r in ch.runs:
            r.font.color.rgb = _rgb(PRIMARY)
        doc.add_paragraph(app.get("comments_prompt",
                                  "Rispondi in chat (le risposte saranno trascritte in feedback/draftN.md)."))

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def has_appendix(docx_path):
    """True if the .docx contains the gate appendix (used to assert final is clean)."""
    from docx import Document
    doc = Document(str(docx_path))
    heads = [p.text.strip() for p in doc.paragraphs]
    return "Domande per continuare" in heads or "Commenti finali" in heads
